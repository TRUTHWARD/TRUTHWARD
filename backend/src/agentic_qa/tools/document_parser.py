# SPDX-License-Identifier: Apache-2.0
from __future__ import annotations

from abc import ABC, abstractmethod
from dataclasses import dataclass
from importlib.util import find_spec
from io import BytesIO
import json
import os
from pathlib import Path, PurePosixPath
import stat
import subprocess
import sys
from typing import Any
from urllib.parse import unquote
import unicodedata
from zipfile import BadZipFile, ZipFile


PDF_MIME_TYPE = "application/pdf"
DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

DOCUMENT_RESOURCE_PROFILE = "document-parser.v1"
MAX_DOCUMENT_BYTES = 1024 * 1024
MAX_DOCUMENT_PAGES = 100
MAX_DOCX_ARCHIVE_ENTRIES = 256
MAX_DOCX_EXPANDED_BYTES = 8 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 200
MAX_DOCUMENT_TEXT_CHARS = 2 * 1024 * 1024
DOCUMENT_PARSE_TIMEOUT_SECONDS = 5.0
DOCUMENT_PARSER_MEMORY_BUDGET_BYTES = 256 * 1024 * 1024
_WORKER_FLAG = "--document-parser-worker"


class DocumentParserError(ValueError):
    """Safe, user-facing failure raised by a controlled document parser adapter."""


@dataclass(frozen=True, slots=True)
class ParsedDocumentText:
    text: str
    parser: str


class DocumentTextParserAdapter(ABC):
    """Tool adapter contract for deterministic text extraction without OCR."""

    @abstractmethod
    def extract_text(self, payload: bytes, *, mime_type: str) -> ParsedDocumentText:
        raise NotImplementedError


class LibraryDocumentTextParserAdapter(DocumentTextParserAdapter):
    """Server-side PDF/DOCX parser executed in a timeout-bounded child process."""

    def __init__(self, *, timeout_seconds: float = DOCUMENT_PARSE_TIMEOUT_SECONDS) -> None:
        if timeout_seconds <= 0:
            raise ValueError("document parser timeout must be positive")
        self.timeout_seconds = float(timeout_seconds)

    def extract_text(self, payload: bytes, *, mime_type: str) -> ParsedDocumentText:
        canonical_mime = (mime_type or "").split(";")[0].strip().lower()
        _preflight_document(payload, canonical_mime)
        _require_parser_dependency(canonical_mime)
        return self._run_worker(payload, canonical_mime)

    def _run_worker(self, payload: bytes, mime_type: str) -> ParsedDocumentText:
        command = [
            sys.executable,
            "-m",
            "agentic_qa.tools.document_parser",
            _WORKER_FLAG,
            mime_type,
            str(MAX_DOCUMENT_PAGES),
            str(MAX_DOCUMENT_TEXT_CHARS),
            str(DOCUMENT_PARSER_MEMORY_BUDGET_BYTES),
        ]
        allowed_environment_names = (
            "LANG",
            "LC_ALL",
            "LD_LIBRARY_PATH",
            "PATH",
            "PATHEXT",
            "PYTHONHOME",
            "SYSTEMROOT",
            "TEMP",
            "TMP",
            "WINDIR",
        )
        environment = {
            name: os.environ[name]
            for name in allowed_environment_names
            if name in os.environ
        }
        import_paths = [value for value in sys.path if value and Path(value).exists()]
        existing_pythonpath = environment.get("PYTHONPATH")
        if existing_pythonpath:
            import_paths.append(existing_pythonpath)
        environment["PYTHONPATH"] = os.pathsep.join(dict.fromkeys(import_paths))
        environment["PYTHONUTF8"] = "1"
        creation_flags = (
            int(getattr(subprocess, "CREATE_NO_WINDOW", 0))
            if os.name == "nt"
            else 0
        )
        try:
            completed = subprocess.run(
                command,
                input=payload,
                stdout=subprocess.PIPE,
                stderr=subprocess.DEVNULL,
                timeout=self.timeout_seconds,
                check=False,
                env=environment,
                creationflags=creation_flags,
            )
        except subprocess.TimeoutExpired as exc:
            raise DocumentParserError(
                f"DOCUMENT_PARSE_TIMEOUT: parser exceeded {self.timeout_seconds:g} second limit"
            ) from exc
        except OSError as exc:
            raise DocumentParserError("DOCUMENT_PARSER_UNAVAILABLE: parser worker could not start") from exc

        if completed.returncode != 0:
            raise DocumentParserError("DOCUMENT_PARSE_FAILED: parser worker failed safely")
        try:
            response = json.loads(completed.stdout.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError, AttributeError) as exc:
            raise DocumentParserError("DOCUMENT_PARSE_FAILED: parser worker returned invalid output") from exc
        if not isinstance(response, dict):
            raise DocumentParserError("DOCUMENT_PARSE_FAILED: parser worker returned invalid output")
        if response.get("status") == "error":
            error = str(response.get("error") or "DOCUMENT_PARSE_FAILED: parser worker failed safely")
            if not _is_safe_parser_error(error):
                error = "DOCUMENT_PARSE_FAILED: parser worker failed safely"
            raise DocumentParserError(error)
        text = response.get("text")
        parser = response.get("parser")
        if not isinstance(text, str) or not isinstance(parser, str):
            raise DocumentParserError("DOCUMENT_PARSE_FAILED: parser worker returned invalid output")
        return ParsedDocumentText(text=text, parser=parser)


def document_text_parser_adapter() -> DocumentTextParserAdapter:
    return LibraryDocumentTextParserAdapter()


def document_resource_limit_snapshot() -> dict[str, int | float | str]:
    return {
        "profile": DOCUMENT_RESOURCE_PROFILE,
        "maxBytes": MAX_DOCUMENT_BYTES,
        "maxPages": MAX_DOCUMENT_PAGES,
        "maxArchiveEntries": MAX_DOCX_ARCHIVE_ENTRIES,
        "maxExpandedBytes": MAX_DOCX_EXPANDED_BYTES,
        "maxCompressionRatio": MAX_DOCX_COMPRESSION_RATIO,
        "maxTextChars": MAX_DOCUMENT_TEXT_CHARS,
        "timeoutSeconds": DOCUMENT_PARSE_TIMEOUT_SECONDS,
        "memoryBudgetBytes": DOCUMENT_PARSER_MEMORY_BUDGET_BYTES,
    }


def _preflight_document(payload: bytes, mime_type: str) -> None:
    if not payload:
        raise DocumentParserError("DOCUMENT_PARSE_FAILED: document is empty")
    if len(payload) > MAX_DOCUMENT_BYTES:
        raise DocumentParserError(
            f"DOCUMENT_RESOURCE_LIMIT: document exceeds {MAX_DOCUMENT_BYTES} byte limit"
        )
    if mime_type == PDF_MIME_TYPE:
        if not payload.startswith(b"%PDF-"):
            raise DocumentParserError("DOCUMENT_PARSE_FAILED: invalid PDF signature")
        return
    if mime_type == DOCX_MIME_TYPE:
        if not payload.startswith(b"PK"):
            raise DocumentParserError("DOCUMENT_PARSE_FAILED: invalid DOCX container")
        _preflight_docx_archive(payload)
        return
    raise DocumentParserError("DOCUMENT_TYPE_UNSUPPORTED: no parser is registered for this mimeType")


def _preflight_docx_archive(payload: bytes) -> None:
    try:
        with ZipFile(BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ARCHIVE_ENTRIES:
                raise DocumentParserError(
                    f"DOCUMENT_RESOURCE_LIMIT: DOCX exceeds {MAX_DOCX_ARCHIVE_ENTRIES} archive entry limit"
                )
            expanded_bytes = 0
            names: set[str] = set()
            for entry in entries:
                original_name = entry.filename
                decoded_name = original_name
                for _ in range(3):
                    next_name = unquote(decoded_name)
                    if next_name == decoded_name:
                        break
                    decoded_name = next_name
                if (
                    "\\" in original_name
                    or decoded_name != original_name
                    or "\x00" in original_name
                    or any(unicodedata.category(character) in {"Cc", "Cf"} for character in original_name)
                    or entry.flag_bits & 0x1
                    or stat.S_ISLNK(entry.external_attr >> 16)
                ):
                    raise DocumentParserError("DOCUMENT_ARCHIVE_INVALID: DOCX contains unsafe archive entries")
                normalized_name = original_name
                path = PurePosixPath(normalized_name)
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or ":" in path.parts[0]
                    or normalized_name in names
                ):
                    raise DocumentParserError("DOCUMENT_ARCHIVE_INVALID: DOCX contains unsafe archive paths")
                names.add(normalized_name)
                expanded_bytes += int(entry.file_size)
                if expanded_bytes > MAX_DOCX_EXPANDED_BYTES:
                    raise DocumentParserError(
                        f"DOCUMENT_RESOURCE_LIMIT: DOCX exceeds {MAX_DOCX_EXPANDED_BYTES} expanded byte limit"
                    )
                if entry.file_size:
                    ratio = entry.file_size / max(1, entry.compress_size)
                    if ratio > MAX_DOCX_COMPRESSION_RATIO:
                        raise DocumentParserError(
                            "DOCUMENT_RESOURCE_LIMIT: DOCX compression ratio exceeds safe limit"
                        )
            required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
            if not required.issubset(names):
                raise DocumentParserError("DOCUMENT_PARSE_FAILED: DOCX package parts are incomplete")
    except DocumentParserError:
        raise
    except (BadZipFile, OSError, ValueError) as exc:
        raise DocumentParserError("DOCUMENT_PARSE_FAILED: malformed DOCX ZIP container") from exc


def _require_parser_dependency(mime_type: str) -> None:
    dependency = "pypdf" if mime_type == PDF_MIME_TYPE else "docx"
    try:
        available = find_spec(dependency) is not None
    except (ImportError, ValueError):
        available = False
    if not available:
        label = "pypdf" if dependency == "pypdf" else "python-docx"
        raise DocumentParserError(
            f"DOCUMENT_PARSER_UNAVAILABLE: {mime_type} parsing requires the declared {label} dependency"
        )


def _extract_document_in_worker(
    payload: bytes,
    mime_type: str,
    *,
    max_pages: int,
    max_text_chars: int,
) -> ParsedDocumentText:
    if mime_type == PDF_MIME_TYPE:
        return _extract_pdf_in_worker(payload, max_pages=max_pages, max_text_chars=max_text_chars)
    if mime_type == DOCX_MIME_TYPE:
        return _extract_docx_in_worker(payload, max_text_chars=max_text_chars)
    raise DocumentParserError("DOCUMENT_TYPE_UNSUPPORTED: no parser is registered for this mimeType")


def _extract_pdf_in_worker(payload: bytes, *, max_pages: int, max_text_chars: int) -> ParsedDocumentText:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentParserError(
            "DOCUMENT_PARSER_UNAVAILABLE: PDF parsing requires the declared pypdf dependency"
        ) from exc
    try:
        reader = PdfReader(BytesIO(payload), strict=False)
        if reader.is_encrypted:
            raise DocumentParserError("DOCUMENT_ENCRYPTED: encrypted PDF is not supported")
        page_count = len(reader.pages)
        if page_count > max_pages:
            raise DocumentParserError(
                f"DOCUMENT_RESOURCE_LIMIT: PDF exceeds {max_pages} page limit"
            )
        parts: list[str] = []
        extracted_chars = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            extracted_chars += len(page_text)
            if extracted_chars > max_text_chars:
                raise DocumentParserError(
                    f"DOCUMENT_RESOURCE_LIMIT: extracted text exceeds {max_text_chars} character limit"
                )
            parts.append(page_text)
    except DocumentParserError:
        raise
    except Exception as exc:
        raise DocumentParserError("DOCUMENT_PARSE_FAILED: PDF text extraction failed") from exc

    text = _normalize_extracted_text("\n\n".join(parts))
    if not text:
        raise DocumentParserError(
            "OCR_REQUIRED: PDF contains no extractable text; scanned PDFs require OCR capability"
        )
    return ParsedDocumentText(text=text, parser="pypdf")


def _extract_docx_in_worker(payload: bytes, *, max_text_chars: int) -> ParsedDocumentText:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentParserError(
            "DOCUMENT_PARSER_UNAVAILABLE: DOCX parsing requires the declared python-docx dependency"
        ) from exc
    try:
        document = Document(BytesIO(payload))
        parts: list[str] = []
        extracted_chars = 0
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                extracted_chars += len(paragraph.text)
                if extracted_chars > max_text_chars:
                    raise DocumentParserError(
                        f"DOCUMENT_RESOURCE_LIMIT: extracted text exceeds {max_text_chars} character limit"
                    )
                parts.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    row_text = "\t".join(values)
                    extracted_chars += len(row_text)
                    if extracted_chars > max_text_chars:
                        raise DocumentParserError(
                            f"DOCUMENT_RESOURCE_LIMIT: extracted text exceeds {max_text_chars} character limit"
                        )
                    parts.append(row_text)
    except DocumentParserError:
        raise
    except Exception as exc:
        raise DocumentParserError("DOCUMENT_PARSE_FAILED: DOCX text extraction failed") from exc

    text = _normalize_extracted_text("\n".join(parts))
    if not text:
        raise DocumentParserError("DOCUMENT_NO_EXTRACTABLE_TEXT: DOCX contains no extractable text")
    return ParsedDocumentText(text=text, parser="python-docx")


def _normalize_extracted_text(value: str) -> str:
    sanitized = "".join(
        character
        for character in value.replace("\r\n", "\n").replace("\r", "\n")
        if character in {"\n", "\t"} or not (ord(character) < 32 or 127 <= ord(character) <= 159)
    )
    return "\n".join(line.rstrip() for line in sanitized.splitlines()).strip()


def _is_safe_parser_error(value: str) -> bool:
    allowed_codes = {
        "DOCUMENT_ARCHIVE_INVALID",
        "DOCUMENT_ENCRYPTED",
        "DOCUMENT_NO_EXTRACTABLE_TEXT",
        "DOCUMENT_PARSE_FAILED",
        "DOCUMENT_PARSER_UNAVAILABLE",
        "DOCUMENT_RESOURCE_LIMIT",
        "DOCUMENT_TYPE_UNSUPPORTED",
        "OCR_REQUIRED",
    }
    code = value.partition(":")[0]
    return code in allowed_codes and len(value) <= 300 and "\x00" not in value


def _apply_worker_memory_limit(memory_budget_bytes: int) -> None:
    if os.name == "nt":
        return
    try:
        import resource

        set_limit = getattr(resource, "setrlimit", None)
        address_space_limit = getattr(resource, "RLIMIT_AS", None)
        if not callable(set_limit) or address_space_limit is None:
            return
        set_limit(address_space_limit, (memory_budget_bytes, memory_budget_bytes))
    except (ImportError, OSError, ValueError):
        return


def _worker_response(result: ParsedDocumentText | None, error: str | None) -> bytes:
    payload: dict[str, Any]
    if error is not None:
        payload = {"status": "error", "error": error}
    elif result is not None:
        payload = {"status": "ok", "text": result.text, "parser": result.parser}
    else:
        payload = {"status": "error", "error": "DOCUMENT_PARSE_FAILED: parser worker failed safely"}
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")


def _run_worker_cli(arguments: list[str]) -> int:
    if len(arguments) != 5 or arguments[0] != _WORKER_FLAG:
        return 2
    mime_type = arguments[1]
    try:
        max_pages = int(arguments[2])
        max_text_chars = int(arguments[3])
        memory_budget_bytes = int(arguments[4])
    except ValueError:
        return 2
    _apply_worker_memory_limit(memory_budget_bytes)
    payload = sys.stdin.buffer.read(MAX_DOCUMENT_BYTES + 1)
    try:
        _preflight_document(payload, mime_type)
        result = _extract_document_in_worker(
            payload,
            mime_type,
            max_pages=max_pages,
            max_text_chars=max_text_chars,
        )
        response = _worker_response(result, None)
    except DocumentParserError as exc:
        response = _worker_response(None, str(exc))
    except Exception:
        response = _worker_response(None, "DOCUMENT_PARSE_FAILED: parser worker failed safely")
    sys.stdout.buffer.write(response)
    return 0


if __name__ == "__main__":
    raise SystemExit(_run_worker_cli(sys.argv[1:]))
